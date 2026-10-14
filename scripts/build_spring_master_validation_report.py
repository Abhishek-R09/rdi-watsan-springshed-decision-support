"""Build a shareable validation report for the processed Spring Master v1."""

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
AUDIT_PATH = ROOT / "data" / "03_outputs" / "validation" / "spring_master_v1_data_quality_audit.json"
OUTPUT_PATH = ROOT / "reports" / "Spring_Master_v1_Data_Quality_Validation_Report.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    run = cell.paragraphs[0].add_run(str(text))
    run.bold = bold
    run.font.size = Pt(8.5)
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, header, bold=True, color=(255, 255, 255))
        shade(cell, "1F4E78")
        if widths:
            cell.width = Inches(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph()
    return table


def add_bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.add_run(text)
    return paragraph


def add_note(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{label}: ")
    r.bold = True
    r.font.color.rgb = RGBColor(31, 78, 120)
    p.add_run(text)


def main():
    audit = json.loads(AUDIT_PATH.read_text(encoding="utf-8"))
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.color.rgb = RGBColor(31, 78, 120)

    title = document.add_heading("Spring Master v1", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Data Quality Validation Report")
    run.bold = True
    run.font.size = Pt(15)
    subtitle.add_run("\nProcessed spring data foundation for SRHU/HIHT RDI WATSAN Project 1")
    subtitle.add_run(f"\nValidation date: {date(2026, 10, 12).strftime('%d %B %Y')}")

    document.add_heading("Validation conclusion", level=1)
    p = document.add_paragraph()
    p.add_run("Decision: ").bold = True
    decision = p.add_run("Suitable for controlled exploratory use and further field validation, with documented caveats.")
    decision.bold = True
    decision.font.color.rgb = RGBColor(31, 100, 60)
    document.add_paragraph(
        "The file is structurally sound: record identifiers are unique, all fields are documented, the master and feature table agree, duplicate-group logic is internally consistent, and the tested numerical and formula relationships have no failures. "
        "It is not a field-verified truth set or a predictive-model training set. Values drawn from map overlays, source documents, or neighbouring colocated records must remain traceable and must not be presented as direct measurements."
    )

    document.add_heading("Data at a glance", level=1)
    add_table(document, ["Measure", "Result", "Quality implication"], [
        ["Source records", "1,842", "Retained as traceable source-level records."],
        ["Primary records", "1,611", "One designated record per proximity group where applicable."],
        ["Fields", "102", "All expected fields are present in the data dictionary."],
        ["Source datasets", "8", "Source dataset is populated for every record."],
        ["Mappable coordinates", "1,803 (98%)", "Spatial overlays available for these records."],
        ["Missing / outside study coverage", "36 missing; 3 outside", "Do not map or derive spatial attributes until resolved."],
        ["Proximity duplicate groups", "217 groups / 448 records", "Candidates for human review, not automatic deletion."],
    ], [1.5, 1.55, 3.85])

    document.add_heading("What was tested", level=1)
    document.add_paragraph("The audit reviewed the processed CSV master and its related processed feature table and data dictionary. It tested:")
    for item in [
        "identity and provenance: blank or repeated record identifiers, blank source attribution, and data-dictionary schema alignment;",
        "cross-file consistency: record identifiers and shared fields between the master and feature table;",
        "geographic quality: coordinate presence, status labels, expected geographic ranges, and decimal precision;",
        "duplicate handling: reported group size and exactly one selected primary record per proximity group;",
        "numeric validity: parseability and broad plausible bounds for terrain, land cover, climate, soil, structure, water quality, discharge and recharge-area fields;",
        "internal relationships: soil texture totals, land-cover component totals, low/high ranges, IMD rainfall ordering, and the rainfall-to-ET0 aridity calculation; and",
        "coverage and provenance flags for the fields most relevant to the three project modules."
    ]:
        add_bullet(document, item)

    document.add_heading("Results by quality parameter", level=1)
    add_table(document, ["Parameter group", "Result", "Assessment / use rule"], [
        ["Identity and traceability", "0 blank IDs; 0 repeated IDs; 0 blank sources", "Pass. Use record_id and source_dataset in all derived outputs."],
        ["Schema and feature-table alignment", "102 expected fields; no missing or extra fields; no shared-field changes", "Pass. Master and processed feature records remain aligned."],
        ["Coordinates", "1,803 valid; 36 missing; 3 outside coverage; 9 low precision", "Conditionally usable. Exclude the 39 records without usable mapped coordinates from spatial analysis."],
        ["Duplicate / colocated records", "217 groups; 0 group-size errors; 0 primary-selection errors", "Pass internally. Confirm whether each group represents the same spring before consolidation."],
        ["Terrain and land cover", "1,803 populated (98%)", "Use as mapped environmental context, not direct field observation."],
        ["Rainfall and ET0", "1,803 populated (98%); 0 aridity-calculation errors", "Use for comparative screening only; spatial climate grids smooth local mountain variation."],
        ["Soil", "1,799 texture records (98%); 0 texture-total failures", "Use as regional overlay. Confirm organic-carbon-density label/units before using it quantitatively."],
        ["Lithology and structural context", "1,761 populated (96%)", "Use as map-scale context; faults, folds and lineaments do not prove a hydraulic connection."],
        ["Spring type", "1,021 populated (55%)", "Insufficient for complete type-based comparisons; retain source text and missingness."],
        ["Discharge", "1,038 populated (56%); 0 low/high ordering errors", "Do not treat source values or colocated fills as current measured discharge without date and method."],
        ["pH / TDS", "702 / 703 populated (38%)", "Insufficient water-quality coverage for safety or trend conclusions."],
        ["Recharge-area range", "356 populated (19%); 0 low/high ordering errors", "Not a training target or proof of a hydrogeological recharge area; retain as source-labelled information only."],
        ["Interventions", "18 populated (1%)", "Too sparse to evaluate intervention patterns or DI impact."],
    ], [1.45, 1.75, 3.7])

    document.add_heading("Detailed validation findings", level=1)
    add_note(document, "Coordinate status", "The status labels correctly identify 36 missing and 3 outside-coverage records. An initial Uttarakhand-only bounding-box screen raised 23 apparent exceptions; review shows these are valid Himmotthan records in Himachal Pradesh. They should remain marked 'ok' if the master is intended to cover the wider Himalayan source area. This audit does not reclassify them as errors.")
    add_note(document, "Elevation comparison", "77 records are flagged as having a source-versus-overlay elevation gap greater than 150 m. This is a reconciliation queue, not a failure: retain both source and mapped elevation, inspect coordinate precision / datum / transcription, and choose a value only for a stated purpose.")
    add_note(document, "No numeric failures", "All tested numeric fields were parseable, within broad physical limits, and internally consistent. This detects structural and arithmetic issues; it does not confirm field accuracy.")
    add_note(document, "Shared-record values", "Some discharge, pH, TDS and recharge-area values were populated from a colocated record. These must preserve their *_from provenance field and should never be silently treated as an independent measurement.")

    document.add_heading("Key numerical checks", level=1)
    add_table(document, ["Check", "Result"], [
        ["Soil clay + silt + sand within 98–102%", "0 exceptions among populated records"],
        ["Tree + cropland + built-up cover above 100%", "0 exceptions"],
        ["Discharge low value greater than high value", "0 exceptions"],
        ["Recharge-area low value greater than high value", "0 exceptions"],
        ["IMD minimum annual rainfall greater than mean", "0 exceptions"],
        ["Aridity index differs from rainfall / ET0 by >0.01", "0 exceptions"],
        ["Out-of-range or unparsable values in 25 tested numeric fields", "0 exceptions"],
    ], [4.3, 2.6])

    document.add_heading("How to use the master safely", level=1)
    document.add_paragraph("The master combines three kinds of information. They must remain distinct in any decision-support output:")
    add_table(document, ["Label", "Examples", "Safe interpretation"], [
        ["Source", "Spring name, source discharge text, water-quality entries, survey details", "Treat as source-attributed information; confirm date, method and site match before use in a field decision."],
        ["Inferred / mapped", "Elevation, slope, land cover, rainfall, soil, GSI geology and structural distances", "Use for first-pass comparison and candidate-area screening only; not as a direct spring measurement."],
        ["Verify", "Duplicate identity, elevation discrepancies, spring type, discharge, water quality, recharge area and DI condition", "Prioritise for WATSAN review, field checking and monitoring."],
    ], [1.1, 2.55, 3.25])

    document.add_heading("Critical limits", level=1)
    for item in audit["known_interpretation_limits"]:
        add_bullet(document, item)
    add_bullet(document, "The 299-point WATSAN data must not be treated as verified ground truth until its provenance, measurement status, labels, coverage and land-use coding have been audited.")
    add_bullet(document, "No before/after DI dataset is present. The file cannot establish DI causality, predict response, or replace rainfall, water-level, discharge/yield and water-quality monitoring.")

    document.add_heading("Recommended next actions", level=1)
    add_table(document, ["Priority", "Action", "Reason"], [
        ["1", "Keep all 39 records without usable map coverage out of spatial scoring; recover or verify coordinates where possible.", "Prevents invalid overlays and false map precision."],
        ["2", "Review the 217 colocated groups with local names, source documents and field GPS; retain both records until confirmed.", "Avoids wrongly merging distinct nearby springs."],
        ["3", "Reconcile the 77 large elevation differences and document the chosen elevation by use case.", "Protects terrain-based screening from coordinate/datum issues."],
        ["4", "Capture field discharge (three bucket-and-stopwatch readings), date/time, season and conditions for priority springs.", "Replaces unverified / shared source values with usable observations."],
        ["5", "Add sample date, method, laboratory/field protocol and re-test status for water quality.", "Current 38% coverage cannot support a safety conclusion."],
        ["6", "Confirm SoilGrids organic-carbon-density definition and units before any calculation using it.", "Avoids a unit-label error in quantitative analysis."],
        ["7", "Use leakage-safe spatial validation before any predictive analysis and exclude recharge-area fields/classes derived from the target.", "Prevents an optimistic but invalid model."],
    ], [0.6, 3.7, 2.6])

    document.add_heading("Audit boundary and reproducibility", level=1)
    document.add_paragraph(
        "This report is a data-quality audit, not hydrogeological validation. It checks file structure, completeness, internal consistency and stated provenance. It does not independently verify original source documents, field measurements, map-overlay accuracy, raster processing choices, hydrogeological connectivity or causality. The accompanying machine-readable audit output records the exact test counts and thresholds so that the checks can be rerun after an update to the master."
    )
    document.add_paragraph("Prepared for the SRHU/HIHT RDI WATSAN Project 1 data foundation. Version reviewed: Spring Master v1.")

    for sec in document.sections:
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("Spring Master v1 Data Quality Validation Report | 12 October 2026")

    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()

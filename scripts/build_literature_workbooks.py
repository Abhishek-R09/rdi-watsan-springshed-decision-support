"""Build the literature and combined-master review workbooks with openpyxl.

Run after stage_watsan_and_spring_atlas.py. This is a Python stand-in for the
literature-related outputs of build_common_spring_workbook.mjs, which needs the
Codex artifact runtime. Values are written as text, exactly as in the CSVs.
"""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
STAGED = ROOT / "data" / "01_staged"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Arial", size=10)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")

LITERATURE_COLUMN_DICTIONARY = [
    ("Column", "Meaning", "Filled for"),
    ("source_reference", "Full citation of the paper or report the row was transcribed from.", "All literature rows"),
    ("value_status_summary", "One-line summary of which values are measured, reported, approximate or estimated. Per-value status is in literature_spring_values_long_v1.csv.", "All literature rows"),
    ("recharge_area_basis_source_value", "How the source defined the area attached to the spring (field hydrogeology, GIS delineation, surface micro-watershed).", "Rathi, Chamba report, Atlas profiles"),
    ("springshed_area_km2_source_value", "Delineated springshed (recharge zone) area in km2, as printed.", "Rathi 2021"),
    ("springshed_elevation_min_m_source_value", "Lowest elevation of the delineated springshed (not the spring outlet).", "Rathi 2021"),
    ("springshed_elevation_max_m_source_value", "Highest elevation of the delineated springshed.", "Rathi 2021"),
    ("spring_group_source_value", "Clubbed springs sharing one springshed, each spring's share of flow, and the group area.", "Rathi 2021"),
    ("catchment_area_ha_source_value", "Surface micro-watershed area in ha. Not a confirmed recharge area.", "Chamba report"),
    ("catchment_elevation_range_m_source_value", "Elevation range of the micro-watershed.", "Chamba report"),
    ("discharge_max_m3_per_day_source_value", "Annual average maximum spring flow in m3/day, as printed (not converted to L/min).", "Rathi 2021"),
    ("discharge_dependable_q75_lpm_source_value", "75% dependable flow by water year, L/min.", "Daniel 2021"),
    ("discharge_annual_series_source_value", "Rainfall and discharge over monitoring years, including before/after recharge works.", "Atlas profiles"),
    ("handpump_discharge_lpm_source_value", "Discharge of the nearby Direct Injection handpump, L/min.", "Chamba report"),
    ("water_quality_other_source_value", "Salinity, chloride, total hardness, iron, nitrate and faecal coliform as printed.", "Atlas profiles"),
    ("recharge_area_slope_aspect_source_value", "Slope and aspect of the recharge area (not the spring point).", "Atlas profiles"),
    ("intervention_source_value", "Recharge or DI works reported for the spring, with dates and quantities where given.", "Chamba report, Daniel 2021, Atlas profiles"),
    ("vulnerability_index_source_value", "Vulnerability index as computed by the source.", "Daniel 2021"),
    ("candidate_match_record_ids", "Existing common-file rows that may be the same spring. A lead to check, not a merge.", "Where a name or coordinate match was seen"),
]

FIELD_MAPPING = [
    ("Common field", "WATSAN source", "Spring Atlas source", "Himmotthan source", "Literature sources", "Rule"),
    ("source_area_name", "Area_name", "", "", "Watershed / tehsil / location text", "Retained as an area label; never assumed to be a spring name."),
    ("spring_name", "", "Spring Name", "Placemark name", "Spring name or number as printed", "Blank where a source does not explicitly provide a spring name."),
    ("latitude_source_value", "Latitude", "Latitude", "Point coordinate", "Printed coordinate (blank in Rathi and Daniel)", "Source text retained."),
    ("longitude_source_value", "Longitude", "Longitude", "Point coordinate", "Printed coordinate (blank in Rathi and Daniel)", "Source text retained."),
    ("elevation_m_source_value", "elevation", "Altitude (m)", "Altitude", "Spring or village elevation as printed", "Source text retained; thousands separators removed for literature rows."),
    ("land_use_source_value", "lulc", "", "Land_use", "", "Blank where source does not supply it."),
    ("slope_aspect_source_value", "aspect", "", "Slope_Aspect", "Micro-watershed aspect (Chamba)", "Blank where source does not supply it."),
    ("slope_degrees_source_value", "slope_deg", "", "Slope", "Micro-watershed average slope (Chamba)", "Blank where source does not supply it."),
    ("lithology_source_value", "", "", "Lithology", "Lithology text", "Blank where source does not supply it."),
    ("spring_type_source_value", "spring_typology", "", "Spring_Type", "Spring type text", "Blank where source does not supply it."),
    ("discharge_lpm_source_value", "discharge_lpm", "", "Discharge (LPM)", "Spot or range discharge in L/min (Chamba, Atlas profiles)", "Source text retained; values may be ranges or other non-numeric entries."),
    ("recharge_area_ha_source_value", "recharge_area_ha", "", "", "Recharge area from field hydrogeology (Atlas profiles only)", "Rathi km2 areas and Chamba micro-watersheds are kept in their own columns."),
]


def read_csv(path):
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.reader(handle))


def add_table_sheet(sheet, rows, widths):
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = WRAP_TOP
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False


def add_records_sheet(sheet, rows):
    for row in rows:
        sheet.append(row)
    sheet.row_dimensions[1].height = 32
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
    for index, name in enumerate(rows[0], start=1):
        width = 16 if index == 1 else 42 if name in {"geometry_wkt", "staging_quality_flag", "source_notes", "source_reference"} else 20
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "B2"
    sheet.sheet_view.showGridLines = False


def add_provenance(sheet, title, count, note):
    values = [
        (title, ""),
        ("Field", "Value"),
        ("Records", count),
        ("Coordinate reference system", "WGS 84 longitude/latitude as supplied by the source or standard KML/PDF convention."),
        ("Source and limitation", note),
        ("Format rule", "All source values remain text so the source representation is not silently changed."),
        ("Not done", "No deduplication, imputation, unit conversion, coordinate correction, or interpretation of source fields."),
    ]
    for row in values:
        sheet.append(row)
    sheet.merge_cells("A1:B1")
    sheet["A1"].fill = HEADER_FILL
    sheet["A1"].font = Font(name="Arial", bold=True, color="FFFFFF", size=14)
    for cell in sheet[2]:
        cell.fill = SUBHEADER_FILL
        cell.font = Font(name="Arial", bold=True, size=10)
    for row in sheet.iter_rows(min_row=3):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = WRAP_TOP
    sheet.column_dimensions["A"].width = 27
    sheet.column_dimensions["B"].width = 100
    sheet.freeze_panes = "A3"
    sheet.sheet_view.showGridLines = False


def build(records_csv, output, title, note, long_csv=None):
    rows = read_csv(records_csv)
    workbook = Workbook()
    records = workbook.active
    records.title = "Spring records"
    add_records_sheet(records, rows)
    add_provenance(workbook.create_sheet("Provenance"), title, len(rows) - 1, note)
    add_table_sheet(workbook.create_sheet("Field mapping"), FIELD_MAPPING, [31, 22, 22, 22, 34, 55])
    add_table_sheet(workbook.create_sheet("Literature columns"), LITERATURE_COLUMN_DICTIONARY, [40, 90, 34])
    if long_csv is not None:
        add_table_sheet(workbook.create_sheet("Value sources"), read_csv(long_csv), [16, 36, 50, 16, 20, 12, 50, 60])
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main():
    literature_dir = STAGED / "literature"
    build(
        literature_dir / "csv" / "literature_spring_records_common_v1.csv",
        literature_dir / "xlsx" / "literature_spring_records_common_v1.xlsx",
        "Literature common spring-record staging extract",
        "Transcribed from Rathi et al. 2021 (SRC-011), Daniel et al. 2021 (SRC-012), the RDI Chamba Spring Assessment Report "
        "(SRC-003) and the Spring Atlas representative spring profiles (SRC-008). Each value's page/table and status is on the Value sources sheet.",
        long_csv=literature_dir / "csv" / "literature_spring_values_long_v1.csv",
    )
    build(
        STAGED / "common" / "csv" / "spring_records_common_v1.csv",
        STAGED / "common" / "xlsx" / "spring_records_common_v1.xlsx",
        "Combined master common spring-record staging extract",
        "Common source-traceable master: WATSAN, Spring Atlas annexure, Himmotthan and literature records. No cross-source "
        "deduplication, coordinate correction, imputation, or inferred matching has been applied; candidate_match_record_ids are leads only.",
    )


if __name__ == "__main__":
    main()

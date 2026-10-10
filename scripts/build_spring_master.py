"""Build the final comprehensive spring master table (one row per common-file record).

Inputs:  data/01_staged/common/csv/spring_records_common_v1.csv
         data/02_processed/spring_features/spring_features_v1.csv
Output:  data/02_processed/spring_master/spring_master_v1.csv (+ data dictionary,
         fill summary) and spring_master_v1.xlsx

How gaps are filled (nothing is invented or interpolated; every filled value says
where it came from in a *_from column):
- Terrain, land cover, rainfall, soil and lineaments are mapped layers, so they are
  given for every spring with a usable coordinate, from the same layer. The source's
  own value is kept beside it for comparison.
- Field-only values (discharge, pH, TDS, spring type, lithology, recharge area) come
  from the record itself. Where the record lacks one, a record from another source
  within 30 m (same colocated group) fills it, named in the *_from column. Otherwise
  the value stays blank and gaps_remaining lists it.
Usage: python scripts/build_spring_master.py
"""

import csv
import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
COMMON_CSV = ROOT / "data" / "01_staged" / "common" / "csv" / "spring_records_common_v1.csv"
FEATURES_CSV = ROOT / "data" / "02_processed" / "spring_features" / "spring_features_v1.csv"
OUT_DIR = ROOT / "data" / "02_processed" / "spring_master"

# Preferred record when several sources list the same spring: richest field data first.
SOURCE_PRIORITY = {"ATP": 0, "RDI": 1, "HIM": 2, "WAT": 3, "DAN": 4, "RAT": 5, "ATL": 6}

SPRING_TYPE_CLASSES = [
    ("karst", "Karst"), ("fractur", "Fracture"), ("fault", "Fracture"), ("joint", "Fracture"),
    ("contact", "Contact"), ("contect", "Contact"), ("depression", "Depression"),
    ("deperession", "Depression"), ("deppression", "Depression"), ("dipression", "Depression"), ("seepage", "Seepage"),
    ("gravity", "Gravity / seepage"), ("weathered", "Fracture"), ("stream", "Stream (not a spring)"),
]
LITHOLOGY_CLASSES = [
    ("dolomit", "Carbonate"), ("limestone", "Carbonate"), ("quartzit", "Quartzite"), ("quarzit", "Quartzite"), ("quatzit", "Quartzite"), ("qaurtzit", "Quartzite"),
    ("phyllit", "Phyllite"), ("schist", "Schist"), ("slate", "Slate"), ("shale", "Shale"),
    ("sandstone", "Sandstone"), ("gneiss", "Gneiss / granite"), ("granite", "Gneiss / granite"),
    ("conglomerate", "Conglomerate"), ("alluvi", "Alluvium / debris"), ("unconsolidated", "Alluvium / debris"), ("metabasic", "Metabasic"), ("debris", "Alluvium / debris"),
    ("meta-volcanic", "Metabasic"), ("metavolcanic", "Metabasic"), ("basic", "Metabasic"), ("gravel", "Alluvium / debris"),
    ("boulder", "Alluvium / debris"), ("sand, silt", "Alluvium / debris"), ("silt, clay", "Alluvium / debris"), ("silt,clay", "Alluvium / debris"),
    ("mudstone", "Shale"), ("claystone", "Shale"), ("siltstone", "Sandstone"), ("marble", "Carbonate"), ("carbonate", "Carbonate"), ("amphibolite", "Metabasic"),
]

FIELD_ONLY = ["discharge_lpm_low", "discharge_lpm_high", "ph", "tds_mg_l", "spring_type_source_text",
              "lithology_source_text", "recharge_area_ha_low", "recharge_area_ha_high"]

# (column, unit, from, meaning)
COLUMNS = [
    ("record_id", "", "Common file", "Record id; prefix shows the source (WAT, HIM, ATL, ATP, RDI, RAT, DAN)."),
    ("source_dataset", "", "Common file", "Dataset the record came from."),
    ("spring_name", "", "Common file", "Spring name as given by the source."),
    ("village", "", "Common file", ""),
    ("block", "", "Common file", ""),
    ("district", "", "Processed features", "District with spellings unified."),
    ("area_name", "", "Common file", "Area / watershed label (WATSAN, literature)."),
    ("latitude", "decimal degrees", "Processed features", ""),
    ("longitude", "decimal degrees", "Processed features", ""),
    ("coordinate_status", "", "Processed features", "ok, missing or outside_study_region. Mapped values exist only for ok."),
    ("coordinate_precision_decimals", "", "Processed features", "2 or fewer decimals means about 1 km uncertainty."),
    ("colocated_group_id", "", "Processed features", "Records within 30 m of each other; probably the same spring listed by two sources."),
    ("colocated_group_size", "", "Processed features", ""),
    ("is_primary_record", "", "Derived", "Y for one record per spring: the only record, or the richest record in its colocated group. Use Y rows to count distinct springs."),
    # terrain
    ("elevation_m", "m", "Copernicus DEM GLO-30", "DEM elevation at the spring (same layer for all springs)."),
    ("elevation_source_m", "m", "Common file", "Elevation the source gave, for comparison."),
    ("elevation_check", "", "Derived", "'gap >150 m' when DEM and source differ by more than 150 m (check the coordinate)."),
    ("slope_deg", "degrees", "Copernicus DEM GLO-30", "Slope at the spring."),
    ("slope_mean_500m_deg", "degrees", "Copernicus DEM GLO-30", "Mean slope within 500 m."),
    ("aspect_class", "", "Copernicus DEM GLO-30", "8-point aspect or Flat."),
    ("relief_1km_m", "m", "Copernicus DEM GLO-30", "Highest minus lowest elevation within 1 km."),
    ("upslope_relief_1km_m", "m", "Copernicus DEM GLO-30", "Highest elevation within 1 km minus the spring elevation."),
    ("tpi_300m_m", "m", "Copernicus DEM GLO-30", "Negative = hollow / valley position."),
    ("slope_source_deg", "degrees", "Common file", "Slope the source gave, for comparison."),
    # land cover
    ("land_cover_point", "", "ESA WorldCover 2021", "Class at the spring (10 m)."),
    ("land_cover_majority_500m", "", "ESA WorldCover 2021", "Most common class within 500 m."),
    ("tree_cover_500m_pct", "%", "ESA WorldCover 2021", ""),
    ("cropland_500m_pct", "%", "ESA WorldCover 2021", ""),
    ("built_up_500m_pct", "%", "ESA WorldCover 2021", ""),
    ("land_use_source_text", "", "Common file", "Land use as the source described it (Himmotthan text; WATSAN code)."),
    # climate
    ("rainfall_mean_annual_mm", "mm/year", "CHIRPS v2 2018-2025 (FAO WaPOR)", "Mean annual rainfall, ~5 km."),
    ("rainfall_cv_pct", "%", "CHIRPS v2 2018-2025", "Year-to-year variation."),
    ("rainfall_monsoon_share_pct", "%", "CHIRPS v2 2018-2025", "Share falling June-September."),
    ("rainfall_dry_season_mm", "mm", "CHIRPS v2 2018-2025", "November-April rain."),
    ("et0_mean_annual_mm", "mm/year", "FAO WaPOR v3", "Reference evapotranspiration."),
    ("aridity_index", "", "Derived", "Rainfall / ET0."),
    ("rainfall_imd_2025_mm", "mm", "IMD 0.25 degree grid", "2025 total; one year only, a cross-check."),
    ("rainfall_source_mm", "mm/year", "Common file", "Rainfall the source gave, for comparison."),
    # soil
    ("soil_clay_pct", "%", "SoilGrids 2.0, 0-5 cm (HWSD v2 0-20 cm where missing)", ""),
    ("soil_silt_pct", "%", "SoilGrids 2.0, 0-5 cm (HWSD v2 where missing)", ""),
    ("soil_sand_pct", "%", "SoilGrids 2.0, 0-5 cm (HWSD v2 where missing)", ""),
    ("soil_texture_from", "", "Derived", "Which layer gave clay/silt/sand."),
    ("soil_bulk_density_g_cm3", "g/cm3", "SoilGrids 2.0, 0-5 cm", ""),
    ("soil_org_carbon_density_kg_m3", "kg/m3", "SoilGrids 2.0, 0-5 cm", ""),
    ("soil_type_wrb", "", "HWSD v2.0", "Dominant soil type (~1 km)."),
    ("soil_drainage", "", "HWSD v2.0", ""),
    ("soil_awc_mm_per_m", "mm/m", "HWSD v2.0", "Available water capacity."),
    # geology
    ("lithology_source_text", "", "Common file / colocated record", "Rock type as described in the field (Himmotthan, Atlas profiles, Chamba report, WATSAN rock type)."),
    ("lithology_class", "", "Derived", "Main rock class from the text: Carbonate, Quartzite, Phyllite, Schist, Slate, ... or Mixed."),
    ("lithology_from", "", "Derived", "Record the lithology text came from."),
    ("geological_structure_text", "", "Common file", "Structure notes (WATSAN, Atlas profiles, Chamba report)."),
    ("lineament_distance_m", "m", "DEM lineament candidates", "Distance to nearest candidate lineament. Unvalidated."),
    ("lineament_density_1km", "km/km2", "DEM lineament candidates", ""),
    ("gsi_lithology", "", "GSI NGDR lithology (Bhukosh)", "Rock description of the GSI map polygon at the spring. Uttarakhand only."),
    ("gsi_lithology_class", "", "Derived", "Main rock class from the GSI description, same classes as lithology_class."),
    ("gsi_group", "", "GSI NGDR lithology", "Stratigraphic group (Garhwal, Jaunsar, Central Crystalline, ...)."),
    ("gsi_formation", "", "GSI NGDR lithology", ""),
    ("gsi_age", "", "GSI NGDR lithology", ""),
    ("gsi_distance_to_contact_m", "m", "GSI NGDR lithology", "Distance to the edge of the spring's map polygon (rock contact)."),
    ("rock_class", "", "Derived", "Best rock class: field description (lithology_class) where there is one, else the GSI map (gsi_lithology_class)."),
    ("rock_class_from", "", "Derived", "field record, or GSI map."),
    ("gsi_fault_distance_m", "m", "GSI NGDR faults (Bhukosh)", "Distance to the nearest mapped fault or thrust."),
    ("gsi_nearest_fault_type", "", "GSI NGDR faults", ""),
    ("gsi_nearest_fault_name", "", "GSI NGDR faults", "Named where GSI names it (MCT, MBT, North Almora Thrust, ...)."),
    ("gsi_thrust_distance_m", "m", "GSI NGDR faults", "Distance to the nearest mapped thrust."),
    ("gsi_fault_density_1km", "km/km2", "GSI NGDR faults", "Mapped fault length within 1 km, per km2."),
    # spring and water
    ("spring_type_source_text", "", "Common file / colocated record", ""),
    ("spring_type_class", "", "Derived", "Fracture, Contact, Depression, Karst, Seepage, Gravity / seepage, or Stream (not a spring)."),
    ("spring_type_from", "", "Derived", ""),
    ("discharge_lpm_low", "L/min", "Common file / colocated record", "Lowest number in the source discharge (Rathi: max flow converted from m3/day)."),
    ("discharge_lpm_high", "L/min", "Common file / colocated record", "Highest number in the source discharge."),
    ("discharge_text", "", "Common file / colocated record", "Discharge as the source wrote it (season, dates)."),
    ("discharge_from", "", "Derived", ""),
    ("ph", "", "Common file / colocated record", ""),
    ("tds_mg_l", "mg/l", "Common file / colocated record", ""),
    ("water_quality_from", "", "Derived", ""),
    ("recharge_area_ha_low", "ha", "Common file / colocated record", "WATSAN estimate, Atlas field delineation, Rathi springshed (km2 x 100)."),
    ("recharge_area_ha_high", "ha", "Common file / colocated record", ""),
    ("recharge_area_from", "", "Derived", ""),
    ("seasonality_text", "", "Common file", "WATSAN seasonality label."),
    ("intervention_text", "", "Common file", "Recharge or DI works reported (Atlas profiles, Chamba, Daniel)."),
    ("vulnerability_index_source", "", "Common file", "Daniel 2021 index (one spring)."),
    # bookkeeping
    ("key_fields_filled", "of 13", "Derived", "How many of 13 key fields have a value: elevation, slope, land cover, rainfall, soil texture, rock type, fault distance, lineament distance, spring type, discharge, pH, TDS, recharge area."),
    ("gaps_remaining", "", "Derived", "Key fields still blank and why."),
    ("notes", "", "Processed features", "Coordinate and elevation notes."),
]
NAMES = [c[0] for c in COLUMNS]


def read(path):
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def classify(text, table, mixed_label=None):
    lowered = (text or "").lower()
    if not lowered.strip():
        return ""
    hits = []
    for key, label in table:
        if key in lowered and label not in hits:
            hits.append(label)
    if not hits:
        return "Other"
    if mixed_label and len(hits) > 1:
        return f"{mixed_label} ({', '.join(hits)})"
    return hits[0]


def field_values(common, feature):
    """Field-only values carried by this record itself."""
    lith = common["lithology_source_value"].strip() or common["rock_type_source_value"].strip()
    max_flow = common["discharge_max_m3_per_day_source_value"]
    max_flow = f"{max_flow} m3/day max flow" if max_flow else ""
    values = {
        "discharge_lpm_low": feature["source_discharge_lpm_low"] or feature["source_max_flow_lpm_converted"],
        "discharge_lpm_high": feature["source_discharge_lpm_high"] or feature["source_max_flow_lpm_converted"],
        "discharge_text": common["discharge_lpm_source_value"] or max_flow,
        "ph": feature["source_ph"],
        "tds_mg_l": feature["source_tds_mg_l"],
        "spring_type_source_text": common["spring_type_source_value"].strip(),
        "lithology_source_text": lith,
        "recharge_area_ha_low": feature["source_recharge_area_ha_low"] or feature["source_springshed_area_ha_converted"],
        "recharge_area_ha_high": feature["source_recharge_area_ha_high"] or feature["source_springshed_area_ha_converted"],
    }
    return {k: (v or "") for k, v in values.items()}


def richness(record_id, values):
    return (sum(1 for k in FIELD_ONLY if values[k]), -SOURCE_PRIORITY.get(record_id[:3], 9))


def main():
    common = {r["record_id"]: r for r in read(COMMON_CSV)}
    features = read(FEATURES_CSV)
    own = {f["record_id"]: field_values(common[f["record_id"]], f) for f in features}
    groups = {}
    for f in features:
        if f["colocated_group_id"]:
            groups.setdefault(f["colocated_group_id"], []).append(f["record_id"])
    primary = {max(members, key=lambda rid: richness(rid, own[rid])) for members in groups.values()}

    rows = []
    for f in features:
        rid = f["record_id"]
        c = common[rid]
        partners = [p for p in groups.get(f["colocated_group_id"], []) if p != rid]
        partners.sort(key=lambda p: richness(p, own[p]), reverse=True)

        def take(keys):
            """Fill a set of keys together from this record, else the richest partner that has them."""
            if any(own[rid][k] for k in keys):
                return {k: own[rid][k] for k in keys}, "this record"
            for p in partners:
                if any(own[p][k] for k in keys):
                    return {k: own[p][k] for k in keys}, f"colocated record {p} (within 30 m)"
            return {k: "" for k in keys}, ""

        row = {n: "" for n in NAMES}
        row.update({
            "record_id": rid, "source_dataset": f["source_dataset"], "spring_name": f["spring_name"], "village": f["village"],
            "block": c["block"], "district": f["district_clean"], "area_name": c["source_area_name"],
            "latitude": f["latitude_dd"], "longitude": f["longitude_dd"], "coordinate_status": f["coordinate_status"],
            "coordinate_precision_decimals": f["coordinate_precision_decimals"],
            "colocated_group_id": f["colocated_group_id"], "colocated_group_size": f["colocated_group_size"],
            "is_primary_record": "Y" if (not f["colocated_group_id"] or rid in primary) else "N",
            "elevation_m": f["dem_elevation_m"], "elevation_source_m": f["source_elevation_m"],
            "elevation_check": "gap >150 m" if f["dem_minus_source_elevation_m"] and abs(float(f["dem_minus_source_elevation_m"])) > 150 else "",
            "slope_deg": f["dem_slope_deg"], "slope_mean_500m_deg": f["dem_mean_slope_500m_deg"], "aspect_class": f["dem_aspect_class"],
            "relief_1km_m": f["dem_relief_1km_m"], "upslope_relief_1km_m": f["dem_upslope_relief_1km_m"], "tpi_300m_m": f["dem_tpi_300m_m"],
            "slope_source_deg": f["source_slope_deg"],
            "land_cover_point": f["wc_point_class"], "land_cover_majority_500m": f["wc_majority_500m_class"],
            "tree_cover_500m_pct": f["wc_share_500m_tree_pct"], "cropland_500m_pct": f["wc_share_500m_cropland_pct"],
            "built_up_500m_pct": f["wc_share_500m_built_pct"], "land_use_source_text": c["land_use_source_value"],
            "rainfall_mean_annual_mm": f["rain_chirps_mean_annual_mm_2018_2025"], "rainfall_cv_pct": f["rain_chirps_cv_annual_pct"],
            "rainfall_monsoon_share_pct": f["rain_chirps_monsoon_share_pct"], "rainfall_dry_season_mm": f["rain_chirps_dry_season_mm"],
            "et0_mean_annual_mm": f["et0_wapor_mean_annual_mm_2018_2025"], "aridity_index": f["aridity_index_p_over_et0"],
            "rainfall_imd_2025_mm": f["rain_imd_annual_mm"], "rainfall_source_mm": f["source_rainfall_mm"],
            "soil_bulk_density_g_cm3": f["sg_bulk_density_0_5cm_g_cm3"], "soil_org_carbon_density_kg_m3": f["sg_org_carbon_density_0_5cm_kg_m3"],
            "soil_type_wrb": f["soil_dominant_wrb"], "soil_drainage": f["soil_drainage"], "soil_awc_mm_per_m": f["soil_awc_mm_per_m"],
            "geological_structure_text": c["geological_structure_source_value"],
            "lineament_distance_m": f["lineament_dem_distance_m"], "lineament_density_1km": f["lineament_dem_density_1km_km_per_km2"],
            "gsi_lithology": f["gsi_lithology"], "gsi_group": f["gsi_group"], "gsi_formation": f["gsi_formation"], "gsi_age": f["gsi_age"],
            "gsi_distance_to_contact_m": f["gsi_distance_to_contact_m"], "gsi_fault_distance_m": f["distance_to_gsi_fault_m"],
            "gsi_nearest_fault_type": f["gsi_nearest_fault_type"], "gsi_nearest_fault_name": f["gsi_nearest_fault_name"],
            "gsi_thrust_distance_m": f["distance_to_gsi_thrust_m"], "gsi_fault_density_1km": f["gsi_fault_density_1km_km_per_km2"],
            "seasonality_text": c["seasonality_source_value"], "intervention_text": c["intervention_source_value"],
            "vulnerability_index_source": c["vulnerability_index_source_value"], "notes": f["processing_notes"],
        })
        if f["sg_clay_0_5cm_pct"]:
            row.update(soil_clay_pct=f["sg_clay_0_5cm_pct"], soil_silt_pct=f["sg_silt_0_5cm_pct"], soil_sand_pct=f["sg_sand_0_5cm_pct"],
                       soil_texture_from="SoilGrids 0-5 cm")
        elif f["soil_topsoil_clay_pct"]:
            row.update(soil_clay_pct=f["soil_topsoil_clay_pct"], soil_silt_pct=f["soil_topsoil_silt_pct"], soil_sand_pct=f["soil_topsoil_sand_pct"],
                       soil_texture_from="HWSD v2 0-20 cm")

        values, source = take(["lithology_source_text"])
        row.update(values, lithology_from=source)
        row["lithology_class"] = classify(values["lithology_source_text"], LITHOLOGY_CLASSES, "Mixed")
        if "ALLUVIUM" in f["gsi_group"].upper():
            row["gsi_lithology_class"] = "Alluvium / debris"
        else:
            row["gsi_lithology_class"] = classify(f["gsi_lithology"], LITHOLOGY_CLASSES, "Mixed")
        if row["lithology_class"]:
            row["rock_class"], row["rock_class_from"] = row["lithology_class"], "field record"
        elif row["gsi_lithology_class"]:
            row["rock_class"], row["rock_class_from"] = row["gsi_lithology_class"], "GSI map"
        values, source = take(["spring_type_source_text"])
        row.update(values, spring_type_from=source)
        row["spring_type_class"] = classify(values["spring_type_source_text"], SPRING_TYPE_CLASSES)
        values, source = take(["discharge_lpm_low", "discharge_lpm_high", "discharge_text"])
        row.update(values, discharge_from=source)
        values, source = take(["ph", "tds_mg_l"])
        row.update(values, water_quality_from=source)
        values, source = take(["recharge_area_ha_low", "recharge_area_ha_high"])
        row.update(values, recharge_area_from=source)

        key = {
            "elevation": row["elevation_m"], "slope": row["slope_deg"], "land cover": row["land_cover_point"],
            "rainfall": row["rainfall_mean_annual_mm"], "soil texture": row["soil_clay_pct"],
            "rock type": row["rock_class"], "fault distance": row["gsi_fault_distance_m"], "lineament distance": row["lineament_distance_m"],
            "spring type": row["spring_type_source_text"], "discharge": row["discharge_lpm_low"], "pH": row["ph"],
            "TDS": row["tds_mg_l"], "recharge area": row["recharge_area_ha_low"],
        }
        row["key_fields_filled"] = sum(1 for v in key.values() if v)
        missing = [k for k, v in key.items() if not v]
        if f["coordinate_status"] != "ok":
            mapped = {"elevation", "slope", "land cover", "rainfall", "soil texture", "lineament distance", "fault distance", "rock type"}
            reason = ["mapped values need a coordinate"] if mapped & set(missing) else []
            missing = [m for m in missing if m not in mapped] + reason
        elif not f["gsi_lithology"]:
            missing = [m + " (GSI layers cover Uttarakhand only)" if m in ("rock type", "fault distance") else m for m in missing]
        row["gaps_remaining"] = "; ".join(missing)
        rows.append(row)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    with (OUT_DIR / "spring_master_v1.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=NAMES, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    with (OUT_DIR / "spring_master_v1_data_dictionary.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["column", "unit", "from", "meaning"])
        writer.writerows(COLUMNS)

    # Fill summary: per key column, filled counts by where the value came from.
    summary = [["column", "filled", "filled_pct", "from_this_record", "from_colocated_record", "blank"]]
    from_cols = {"lithology_source_text": "lithology_from", "spring_type_source_text": "spring_type_from",
                 "discharge_lpm_low": "discharge_from", "ph": "water_quality_from", "tds_mg_l": "water_quality_from",
                 "recharge_area_ha_low": "recharge_area_from"}
    for name in NAMES:
        filled = sum(1 for r in rows if r[name] not in ("", None))
        this = other = ""
        if name in from_cols:
            this = sum(1 for r in rows if r[name] and r[from_cols[name]] == "this record")
            other = sum(1 for r in rows if r[name] and r[from_cols[name]].startswith("colocated"))
        summary.append([name, filled, f"{100 * filled / len(rows):.0f}", this, other, len(rows) - filled])
    with (OUT_DIR / "spring_master_v1_fill_summary.csv").open("w", newline="", encoding="utf-8") as handle:
        csv.writer(handle, lineterminator="\n").writerows(summary)

    write_workbook(rows, summary)
    primaries = sum(1 for r in rows if r["is_primary_record"] == "Y")
    print(f"{len(rows)} records, {primaries} primary (distinct springs)")
    print(Counter(r["key_fields_filled"] for r in rows if r["is_primary_record"] == "Y"))


GUIDE = [
    ("What this is", "One row per spring record from the common file (WATSAN, Spring Atlas annexure, Himmotthan, Atlas profiles, "
                     "Chamba report, Rathi 2021, Daniel 2021), with every gap filled that public data or a matching record can fill."),
    ("Counting springs", "Some springs are listed by two sources. Filter is_primary_record = Y to keep one row per spring."),
    ("Mapped values", "Terrain, land cover, rainfall, soil, GSI geology, faults and lineaments come from the same layer for every spring, so springs can be "
                      "compared fairly. The source's own value is kept in a *_source column beside it."),
    ("Field values", "Discharge, pH, TDS, spring type, lithology and recharge area come from the record itself. Where it had none, a "
                     "record from another source within 30 m filled it; the *_from column names that record."),
    ("Not filled", "Nothing is estimated, averaged from neighbours or invented. Blank means no source has it; gaps_remaining lists what is missing."),
    ("Caveats", "Layer values describe the spring point or a circle around it, not the real recharge area. CHIRPS rainfall is ~5 km, "
                "HWSD soil ~1 km, SoilGrids 250 m. DEM lineaments are unvalidated candidates. IMD rainfall is 2025 only. "
                "GSI geology is a state-scale map: a spring's polygon is a map overlay, not a field check, and it covers Uttarakhand only."),
    ("Rebuild", "python scripts/build_spring_features.py ... then python scripts/build_spring_master.py. "
                "See docs/data-gap-analysis.md."),
]


def write_workbook(rows, summary):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    body_font = Font(name="Arial", size=10)
    numeric = re.compile(r"^-?\d+(\.\d+)?$")
    wb = Workbook()
    guide = wb.active
    guide.title = "Guide"
    guide.append(["Spring master v1", ""])
    guide["A1"].font = Font(name="Arial", bold=True, size=14)
    for item in GUIDE:
        guide.append(list(item))
    for row in guide.iter_rows(min_row=2):
        row[0].font = Font(name="Arial", bold=True, size=10)
        row[1].font = body_font
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[0].alignment = Alignment(vertical="top")
    guide.column_dimensions["A"].width = 20
    guide.column_dimensions["B"].width = 110

    def table(sheet, header, body, widths=None):
        sheet.append(header)
        for values in body:
            sheet.append([float(v) if isinstance(v, str) and numeric.match(v) else v for v in values])
        for cell in sheet[1]:
            cell.fill, cell.font = header_fill, header_font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
        for i, name in enumerate(header, start=1):
            sheet.column_dimensions[get_column_letter(i)].width = (widths or {}).get(name, 16)
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions

    wide = {n: 40 for n in ("lithology_source_text", "spring_type_source_text", "discharge_text", "gaps_remaining", "notes",
                            "geological_structure_text", "intervention_text", "land_use_source_text")}
    wide.update({n: 30 for n in ("lithology_from", "spring_type_from", "discharge_from", "water_quality_from", "recharge_area_from", "source_dataset")})
    table(wb.create_sheet("Springs"), NAMES, [[r[n] for n in NAMES] for r in rows], wide)
    table(wb.create_sheet("Data dictionary"), ["column", "unit", "from", "meaning"], [list(c) for c in COLUMNS],
          {"column": 30, "from": 40, "meaning": 90})
    table(wb.create_sheet("Fill summary"), summary[0], summary[1:], {"column": 32})
    wb.save(OUT_DIR / "spring_master_v1.xlsx")


if __name__ == "__main__":
    main()
